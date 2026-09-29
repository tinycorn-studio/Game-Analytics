import os
import pandas as pd
from typing import List, Dict, Any

class ReportGenerator:
    @staticmethod
    def generate_reports(levels: List[Dict[str, Any]], project_dir: str):
        """
        Generates both report.xlsx and report.csv in the specified project directory.
        """
        rows = []
        for lvl in levels:
            rows.append({
                "Level": f"Level {lvl['level']:02d}",
                "Thời gian bắt đầu": lvl.get("start_time_str", ""),
                "Thời gian kết thúc": lvl.get("end_time_str", ""),
                "Thời lượng giải (giây)": lvl.get("duration_seconds", 0),
                "Ảnh bắt đầu màn": lvl.get("board_image_rel", ""),
                "Ảnh chiến thắng": lvl.get("victory_image_rel", ""),
                "Trạng thái": lvl.get("status", "Completed")
            })

        df = pd.DataFrame(rows)
        
        # 1. Export CSV
        csv_path = os.path.join(project_dir, "report.csv")
        df.to_csv(csv_path, index=False, encoding="utf-8-sig")
        
        # 2. Export Excel with OpenPyXL formatting
        xlsx_path = os.path.join(project_dir, "report.xlsx")
        with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Level_Matrix")
            ws = writer.sheets["Level_Matrix"]
            
            # Format header row
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='E5E7EB'),
                right=Side(style='thin', color='E5E7EB'),
                top=Side(style='thin', color='E5E7EB'),
                bottom=Side(style='thin', color='E5E7EB')
            )
            
            for col_num, col_title in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                
            # Auto-fit column widths & center text
            for row in ws.iter_rows(min_row=2, max_row=len(df)+1, min_col=1, max_col=len(df.columns)):
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border
                    
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
                
        return xlsx_path, csv_path
